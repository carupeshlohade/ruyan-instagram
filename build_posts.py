"""Reads the October calendar workbook, renders 1080x1350 branded images for Static/Carousel rows,
and writes queue.csv (Reels get status needs_video). Re-runnable. Usage: python build_posts.py CALENDAR.xlsx"""
import sys, os, re, textwrap, datetime as dt
import openpyxl
from PIL import Image, ImageDraw, ImageFont
from common import write_queue

BG, FG, ACC, MUT = (17,17,17), (245,242,234), (217,164,65), (150,146,138)
W, H = 1080, 1350
F = "/usr/share/fonts/truetype/google-fonts/Poppins-%s.ttf"
def font(w, s): return ImageFont.truetype(F % w, s)

def wrap(d, text, fnt, maxw):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if d.textlength(t, font=fnt) <= maxw: cur = t
        else: lines.append(cur); cur = w
    return lines + ([cur] if cur else [])

def fit(d, text, weight, maxw, maxh, start, minsize=34):
    for s in range(start, minsize - 1, -2):
        f = font(weight, s); L = wrap(d, text, f, maxw)
        if len(L) * s * 1.28 <= maxh: return f, L, s
    f = font(weight, minsize); return f, wrap(d, text, f, maxw), minsize

def frame(tag, n=None, total=None):
    im = Image.new("RGB", (W, H), BG); d = ImageDraw.Draw(im)
    d.ellipse((70, 70, 150, 150), outline=FG, width=4)
    d.text((110, 110), "R", font=font("Bold", 46), fill=FG, anchor="mm")
    d.text((172, 110), "RUYAN CORPORATE SERVICES", font=font("Bold", 24), fill=FG, anchor="lm")
    d.text((70, 215), tag.upper(), font=font("Bold", 26), fill=ACC)
    if n: d.text((W - 70, 110), f"{n}/{total}", font=font("Medium", 26), fill=MUT, anchor="rm")
    d.line((70, H - 120, W - 70, H - 120), fill=(60, 60, 60), width=2)
    d.text((70, H - 80), "Accounting | GST | Tax | Loans | Subsidy | ROC", font=font("Regular", 24), fill=MUT, anchor="lm")
    return im, d

def body(d, text, weight, start, top=290, bottom=H - 170, color=FG):
    f, L, s = fit(d, text, weight, W - 140, bottom - top, start)
    y = top
    for ln in L: d.text((70, y), ln, font=f, fill=color); y += int(s * 1.28)
    return y

def caption(r, hook, pts, cta, tags):
    parts = [hook, ""]
    if r["Format"] == "Carousel": parts += ["Swipe through to the end, and save it for later.", ""]
    elif r["Format"] == "Static": parts += [p for p in pts[:4]] + [""]
    parts += [cta, "", "Enquiries: send us a message or call 7020570215.", "General information only. Please consult a professional for your specific case.", "", tags]
    return "\n".join(parts)[:2150]

def main(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True); ws = wb["Oct 2026 Calendar"]
    hdr = [c.value for c in ws[4]]; rows = []
    for raw in ws.iter_rows(min_row=5, values_only=True):
        if not raw[0]: continue
        r = dict(zip(hdr, raw)); date = r["Date"].date()
        hook, cta = r["Hook (spoken / headline)"], r["CTA"]
        pts = [re.sub(r"^(one date card|checklist frame):\s*", "", p.strip(), flags=re.I) for p in str(r["Key points / on-screen text"]).split("|") if p.strip()]
        pts = [p for p in pts if not re.match(r"^(slide \d+|cta slide|close:?)\b", p, re.I)]
        pid = f"{date.isoformat()}-{r['Format'].lower()}"
        files = []
        if r["Format"] == "Static":
            im, d = frame(r["Pillar"]); y = body(d, hook, "Bold", 70, bottom=700)
            yy = y + 50
            for p in pts[:4]:
                fp, Lp, sp = fit(d, "- " + p, "Regular", W - 140, 120, 36)
                for ln in Lp:
                    if yy < H - 190: d.text((70, yy), ln, font=fp, fill=MUT); yy += int(sp * 1.3)
                yy += 14
            fn = pid + ".png"; im.save("media/" + fn); files = [fn]
        elif r["Format"] == "Carousel":
            slides = [("Swipe", hook)] + [(f"Point {i+1}", p) for i, p in enumerate(pts[:8])] + [("Next step", cta + "  Call 7020570215")]
            for i, (tag, txt) in enumerate(slides, 1):
                im, d = frame(r["Pillar"] if i == 1 else tag, i, len(slides))
                body(d, txt, "Bold" if i in (1, len(slides)) else "Bold", 66 if i == 1 else 56)
                fn = f"{pid}-{i}.png"; im.save("media/" + fn); files.append(fn)
        t = r["Post time (IST)"]; t = t.strftime("%H:%M") if hasattr(t, "strftime") else "19:30"
        rows.append({"id": pid, "date": date.isoformat(), "time": t, "format": r["Format"],
                     "media": ";".join(files), "caption": caption(r, hook, pts, cta, r["Hashtags"]),
                     "status": "needs_video" if r["Format"] == "Reel" else "draft",
                     "posted_media_id": "", "posted_at": "", "error": ""})
    write_queue(rows); print(len(rows), "rows;", len(os.listdir("media")), "media files")

if __name__ == "__main__": main(sys.argv[1])
